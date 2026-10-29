"""Recalculates the workbook formulas and compares them with stats computed
independently from the raw answers, so a broken formula cannot pass silently."""
import json
import sys

import formulas
import openpyxl

from build_workbook import CHANGE_MAX, CHANGE_MIN, HIST_FIRST, NBINS, NQ, OUT, WEB_DATA
from generate_data import LETTERS, QUESTIONS, SCENARIOS, generate

# The web dashboard must read exactly the records that sit in the workbook sheets.
web = json.loads((WEB_DATA / "dashboard-data.json").read_text())
book = openpyxl.load_workbook(OUT, read_only=True)
mismatches = 0
for key, scenario in web["scenarios"].items():
    for stage in ("pre", "post"):
        sheet = book[f"{key}_{stage.capitalize()}"]
        rows = [[r[0], r[3], r[5]] for r in sheet.iter_rows(min_row=3, values_only=True)]
        mismatches += rows != scenario[stage]
        names = {r[0]: r[12] for r in sheet.iter_rows(min_row=3, values_only=True)}
        mismatches += names != {p["id"]: p["name"] for p in scenario["participants"]}
        print("ok  " if rows == scenario[stage] else "FAIL", f"{key} {stage}: dashboard JSON == workbook sheet ({len(rows)} rows)")
if mismatches:
    sys.exit("dashboard data differs from workbook")

xl = formulas.ExcelModel().loads(str(OUT)).finish()
sol = xl.calculate()
get = lambda sheet, cell: sol[f"'[{OUT.name}]{sheet.upper()}'!{cell}"].value[0, 0]

failures = 0
def check(label, got, want, tol=1e-9):
    global failures
    ok = abs(float(got) - want) <= tol
    failures += not ok
    print(("ok  " if ok else "FAIL"), label, got, want)

for key in SCENARIOS:
    data = generate(key)
    def score(stage):
        form = "a" if stage == "Pre" else "b"
        per = {}
        for pid, qid, letter in data[stage]:
            q = next(x for x in QUESTIONS if x["id"] == qid)
            per[pid] = per.get(pid, 0) + (LETTERS[q[form]["answer"]] == letter)
        return per
    pre, post = score("Pre"), score("Post")
    n = len(pre)
    avg_pre, avg_post = sum(pre.values()) / n / NQ, sum(post.values()) / n / NQ
    sheet = f"{key}_Analysis"
    check(f"{key} n", get(sheet, "B4"), n)
    check(f"{key} avg pre %", get(sheet, "B7"), avg_pre)
    check(f"{key} avg post %", get(sheet, "B8"), avg_post)
    check(f"{key} abs pp", get(sheet, "B9"), (avg_post - avg_pre) * 100)
    check(f"{key} relative", get(sheet, "B10"), (avg_post - avg_pre) / avg_pre)
    check(f"{key} improved share", get(sheet, "B11"), sum(post[p] > pre[p] for p in pre) / n)
    d = [post[p] - pre[p] for p in pre]
    assert min(d) >= CHANGE_MIN and max(d) <= CHANGE_MAX, f"changes {min(d)}..{max(d)} outside histogram bins"
    check(f"{key} categories sum", sum(get(sheet, f"C{r}") for r in range(23, 27)), n)
    check(f"{key} substantial", get(sheet, "C23"), sum(x >= 3 for x in d))
    check(f"{key} lower", get(sheet, "C26"), sum(x < 0 for x in d))
    check(f"{key} histogram sum", sum(get(sheet, f"B{r}") for r in range(HIST_FIRST, HIST_FIRST + NBINS)), n)
print("FAILURES:", failures)
sys.exit(failures)
